"""Export OSS statistics as a filterable Excel workbook."""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, Mapping, Optional, Sequence

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet


TITLE_FILL = PatternFill("solid", fgColor="24313C")
SECTION_FILL = PatternFill("solid", fgColor="137A63")
NOTE_FILL = PatternFill("solid", fgColor="F2F4F5")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="B7C3BE"))
PERCENT_FORMAT = "0.0%"
DECIMAL_FORMAT = "0.0"


def export_xlsx(report: Dict[str, Any], output: Path) -> Path:
    """Write an OSS statistics report to an xlsx file and return its path."""
    workbook = Workbook()
    workbook.properties.creator = "OpenXiangShan"
    workbook.properties.title = (
        f"{report['organization']} {report['period']['label']} 开源生态质量月报"
    )
    workbook.properties.subject = "GitHub 开源生态与 GitHub Actions 月度统计"
    workbook.properties.description = "由 dingtalk_robot.oss_stats 自动生成"
    summary_sheet = workbook.active
    summary_sheet.title = "组织汇总"
    _write_summary(summary_sheet, report)
    _write_activity_sheet(workbook.create_sheet("仓库活跃度"), report)
    _write_collaboration_sheet(workbook.create_sheet("协作质量"), report)
    _write_actions_sheet(workbook.create_sheet("Actions质量"), report)
    _write_ci_coverage_sheet(workbook.create_sheet("CI覆盖"), report)
    _write_workflows_sheet(workbook.create_sheet("工作流明细"), report)
    _write_metadata_sheet(workbook.create_sheet("说明"), report)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def _write_summary(sheet: Worksheet, report: Dict[str, Any]) -> None:
    summary = report["summary"]
    vitality = summary["community_vitality"]
    developers = summary["developer_base"]
    collaboration = summary["collaboration_efficiency"]
    actions = summary["continuous_integration"]
    coverage = report["coverage"]
    ci_coverage = coverage["github_actions"]
    period = report["period"]

    _write_title(
        sheet,
        f"{report['organization']} {period['label']} 开源生态质量月报",
        _period_description(period),
        6,
    )
    row = 4
    row = _write_summary_table(
        sheet,
        row,
        "范围与当前规模",
        ["指标", "数值", "说明"],
        [
            ["公开仓库", report["repository_count"], "本次采集的公开仓库"],
            [
                "组织汇总仓库",
                report["aggregation"]["repository_count"],
                "默认排除 Fork 仓库，避免重复统计上游活动",
            ],
            ["当前 Stars", summary["current"]["stars"], "采集时点累计值"],
            ["当前 Forks", summary["current"]["forks"], "采集时点累计值"],
            [
                "仓库历史完整覆盖",
                coverage["complete_repositories"] / coverage["total_repositories"]
                if coverage["total_repositories"]
                else None,
                f"{coverage['complete_repositories']}/{coverage['total_repositories']}",
            ],
        ],
        "SummaryScope",
    )
    sheet.cell(row - 2, 2).number_format = PERCENT_FORMAT
    row = _write_summary_table(
        sheet,
        row,
        "月度活跃度",
        ["指标", "数值", "单位/口径"],
        [
            ["新增 Stars", "N/A", "GitHub 访问限制"],
            ["新增 Forks", vitality["forks_added"], "个"],
            ["默认分支 Commits", vitality["commits"], "次"],
            ["代码新增", vitality["additions"], "行"],
            ["代码删除", vitality["deletions"], "行"],
            ["代码变更", vitality["lines_changed"], "行"],
            ["新建 Issue", vitality["issues_opened"], "个"],
            ["Issue 评论", vitality["issue_comments"], "次"],
            ["新建 PR", vitality["prs_opened"], "个"],
            ["PR 评论", vitality["pr_comments"], "次"],
            ["合并 PR", collaboration["prs_merged"], "个"],
            ["PR 评审", collaboration["pr_reviews"], "次"],
        ],
        "SummaryVitality",
    )
    row = _write_summary_table(
        sheet,
        row,
        "贡献者",
        ["指标", "人数"],
        [
            ["活跃贡献者", developers["contributors"]],
            ["代码贡献者", developers["code_contributors"]],
            ["非代码贡献者", developers["non_code_contributors"]],
        ],
        "SummaryDevelopers",
    )
    row = _write_summary_table(
        sheet,
        row,
        "协作质量",
        ["对象", "未响应率", "首次响应(h)", "处理时长(h)"],
        [
            [
                "Issue",
                collaboration["issue_unresponsive_rate"],
                collaboration["issue_first_response_hours"],
                collaboration["issue_resolution_hours"],
            ],
            [
                "PR",
                collaboration["pr_unresponsive_rate"],
                collaboration["pr_first_response_hours"],
                collaboration["pr_resolution_hours"],
            ],
        ],
        "SummaryCollaboration",
        {1: PERCENT_FORMAT},
    )
    row = _write_summary_table(
        sheet,
        row,
        "PR 协作质量",
        ["合并率", "评审参与率", "平均交互", "PR/Issue 关联率"],
        [[
            collaboration["pr_merge_rate"],
            collaboration["pr_review_participation_rate"],
            collaboration["pr_average_interactions"],
            collaboration["pr_issue_link_rate"],
        ]],
        "SummaryPullRequests",
        {0: PERCENT_FORMAT, 1: PERCENT_FORMAT, 3: PERCENT_FORMAT},
    )
    row = _write_summary_table(
        sheet,
        row,
        "GitHub Actions 运行结果",
        ["Runs", "成功", "失败", "取消", "跳过", "需处理"],
        [[
            actions["runs"],
            actions["successful_runs"],
            actions["conclusions"].get("failure", 0),
            actions["conclusions"].get("cancelled", 0),
            actions["conclusions"].get("skipped", 0),
            actions["conclusions"].get("action_required", 0),
        ]],
        "SummaryActionsRuns",
    )
    row = _write_summary_table(
        sheet,
        row,
        "GitHub Actions 稳定性",
        ["成功率", "首次通过率", "重跑", "重跑率", "重跑后成功"],
        [[
            actions["success_rate"],
            actions["first_attempt_pass_rate"],
            actions["rerun_runs"],
            actions["rerun_rate"],
            actions["successful_after_rerun"],
        ]],
        "SummaryActionsReliability",
        {0: PERCENT_FORMAT, 1: PERCENT_FORMAT, 3: PERCENT_FORMAT},
    )
    row = _write_summary_table(
        sheet,
        row,
        "GitHub Actions 时长",
        ["总时长(h)", "平均(min)", "P50(min)", "P90(min)", "时长样本"],
        [[
            _hours_from_minutes(actions["duration_minutes"]["total"]),
            actions["duration_minutes"]["average"],
            actions["duration_minutes"]["p50"],
            actions["duration_minutes"]["p90"],
            actions["duration_minutes"]["samples"],
        ]],
        "SummaryActionsDuration",
        {0: DECIMAL_FORMAT, 1: DECIMAL_FORMAT, 2: DECIMAL_FORMAT, 3: DECIMAL_FORMAT},
    )
    row = _write_summary_table(
        sheet,
        row,
        "GitHub Actions 排队时长",
        ["总计(min)", "平均(min)", "P50(min)", "P90(min)", "排队样本"],
        [[
            actions["queue_minutes"]["total"],
            actions["queue_minutes"]["average"],
            actions["queue_minutes"]["p50"],
            actions["queue_minutes"]["p90"],
            actions["queue_minutes"]["samples"],
        ]],
        "SummaryActionsQueue",
        {0: DECIMAL_FORMAT, 1: DECIMAL_FORMAT, 2: DECIMAL_FORMAT, 3: DECIMAL_FORMAT},
    )
    _write_summary_table(
        sheet,
        row,
        "CI 覆盖",
        ["Workflow", "有运行仓库", "仓库覆盖率", "PR 覆盖率", "默认分支 Commit 覆盖率"],
        [[
            actions["workflow_count"],
            f"{ci_coverage['repositories_with_runs']}/{ci_coverage['aggregate_repositories']}",
            ci_coverage["repository_run_coverage_rate"],
            actions["coverage"]["pull_request_rate"],
            actions["coverage"]["default_branch_commit_rate"],
        ]],
        "SummaryCiCoverage",
        {2: PERCENT_FORMAT, 3: PERCENT_FORMAT, 4: PERCENT_FORMAT},
    )
    sheet.freeze_panes = "A4"
    _finish_sheet(sheet, max_width=48)


def _write_activity_sheet(sheet: Worksheet, report: Dict[str, Any]) -> None:
    headers = [
        "排名",
        "仓库",
        "完整名称",
        "Fork",
        "归档",
        "计入组织汇总",
        "活跃度分",
        "当前 Stars",
        "当前 Forks",
        "新增 Forks",
        "Commits",
        "新增行",
        "删除行",
        "变更行",
        "新建 Issue",
        "Issue 评论",
        "新建 PR",
        "PR 评论",
        "合并 PR",
        "活跃贡献者",
        "代码贡献者",
        "非代码贡献者",
    ]
    include_forks = report["aggregation"]["includes_fork_repositories"]
    rows = []
    for rank, repo in enumerate(report["repositories"], 1):
        vitality = repo["community_vitality"]
        rows.append(
            [
                rank,
                repo["name"],
                repo["full_name"],
                _yes_no(repo["fork"]),
                _yes_no(repo["archived"]),
                _yes_no(include_forks or not repo["fork"]),
                repo["activity_score"],
                repo["current"]["stars"],
                repo["current"]["forks"],
                vitality["forks_added"],
                vitality["commits"],
                vitality["additions"],
                vitality["deletions"],
                vitality["lines_changed"],
                vitality["issues_opened"],
                vitality["issue_comments"],
                vitality["prs_opened"],
                vitality["pr_comments"],
                repo["collaboration_efficiency"]["prs_merged"],
                repo["developer_base"]["contributors"],
                repo["developer_base"]["code_contributors"],
                repo["developer_base"]["non_code_contributors"],
            ]
        )
    _write_data_sheet(sheet, "仓库活跃度", report, headers, rows, "RepositoryActivity")
    _add_repository_links(sheet, report["repositories"], column=2)


def _write_collaboration_sheet(sheet: Worksheet, report: Dict[str, Any]) -> None:
    headers = [
        "仓库",
        "Issue 关闭",
        "Issue 未响应率",
        "Issue 首次响应(h)",
        "Issue 处理时长(h)",
        "PR 关闭",
        "PR 合并",
        "PR 合并率",
        "PR 未响应率",
        "PR 首次响应(h)",
        "PR 处理时长(h)",
        "PR/Issue 关联率",
        "评审参与率",
        "平均交互",
        "PR 评审",
    ]
    rows = []
    for repo in report["repositories"]:
        item = repo["collaboration_efficiency"]
        rows.append(
            [
                repo["name"],
                item["issues_closed"],
                item["issue_unresponsive_rate"],
                item["issue_first_response_hours"],
                item["issue_resolution_hours"],
                item["prs_closed"],
                item["prs_merged"],
                item["pr_merge_rate"],
                item["pr_unresponsive_rate"],
                item["pr_first_response_hours"],
                item["pr_resolution_hours"],
                item["pr_issue_link_rate"],
                item["pr_review_participation_rate"],
                item["pr_average_interactions"],
                item["pr_reviews"],
            ]
        )
    _write_data_sheet(
        sheet,
        "仓库协作质量",
        report,
        headers,
        rows,
        "RepositoryCollaboration",
        percent_columns={3, 8, 9, 12, 13},
        decimal_columns={4, 5, 10, 11, 14},
    )
    _add_repository_links(sheet, report["repositories"], column=1)


def _write_actions_sheet(sheet: Worksheet, report: Dict[str, Any]) -> None:
    headers = [
        "仓库",
        "状态",
        "Runs",
        "已完成",
        "成功",
        "失败",
        "取消",
        "跳过",
        "需处理",
        "成功率",
        "首次通过率",
        "重跑",
        "重跑率",
        "重跑后成功",
        "Workflow",
        "总时长(h)",
        "平均时长(min)",
        "P50(min)",
        "P90(min)",
        "平均排队(min)",
        "P50 排队(min)",
        "P90 排队(min)",
    ]
    rows = []
    for repo in report["repositories"]:
        item = repo["continuous_integration"]
        rows.append(
            [
                repo["name"],
                item["status"],
                item["runs"],
                item["completed_runs"],
                item["successful_runs"],
                item["conclusions"].get("failure", 0),
                item["conclusions"].get("cancelled", 0),
                item["conclusions"].get("skipped", 0),
                item["conclusions"].get("action_required", 0),
                item["success_rate"],
                item["first_attempt_pass_rate"],
                item["rerun_runs"],
                item["rerun_rate"],
                item["successful_after_rerun"],
                item["workflow_count"],
                _hours_from_minutes(item["duration_minutes"]["total"]),
                item["duration_minutes"]["average"],
                item["duration_minutes"]["p50"],
                item["duration_minutes"]["p90"],
                item["queue_minutes"]["average"],
                item["queue_minutes"]["p50"],
                item["queue_minutes"]["p90"],
            ]
        )
    _write_data_sheet(
        sheet,
        "GitHub Actions 质量",
        report,
        headers,
        rows,
        "RepositoryActions",
        percent_columns={10, 11, 13},
        decimal_columns={16, 17, 18, 19, 20, 21, 22},
    )
    _add_repository_links(sheet, report["repositories"], column=1)


def _write_ci_coverage_sheet(sheet: Worksheet, report: Dict[str, Any]) -> None:
    headers = [
        "仓库",
        "有 Actions 运行",
        "Workflow",
        "PR",
        "有 Actions 的 PR",
        "PR 覆盖率",
        "默认分支 Commits",
        "有 Actions 的 Commits",
        "Commit 覆盖率",
        "采集请求",
        "区间分片",
        "采集错误",
    ]
    rows = []
    for repo in report["repositories"]:
        actions = repo["continuous_integration"]
        coverage = actions["coverage"]
        collection = actions["collection"]
        rows.append(
            [
                repo["name"],
                _yes_no(coverage["has_runs"]),
                actions["workflow_count"],
                coverage["pull_requests"],
                coverage["pull_requests_with_runs"],
                coverage["pull_request_rate"],
                coverage["default_branch_commits"],
                coverage["default_branch_commits_with_runs"],
                coverage["default_branch_commit_rate"],
                collection["requests"],
                collection["partitions"],
                collection["error"] or "",
            ]
        )
    _write_data_sheet(
        sheet,
        "GitHub Actions 覆盖",
        report,
        headers,
        rows,
        "RepositoryCiCoverage",
        percent_columns={6, 9},
    )
    _add_repository_links(sheet, report["repositories"], column=1)


def _write_workflows_sheet(sheet: Worksheet, report: Dict[str, Any]) -> None:
    headers = ["仓库", "Workflow ID", "名称", "配置路径", "Runs"]
    rows = []
    repository_urls = {
        repo["full_name"]: repo["url"] for repo in report["repositories"]
    }
    links = []
    for repo in report["repositories"]:
        for workflow in repo["continuous_integration"]["workflows"]:
            rows.append(
                [
                    workflow["repository"],
                    workflow["workflow_id"],
                    workflow["name"],
                    workflow["path"],
                    workflow["runs"],
                ]
            )
            links.append(repository_urls.get(workflow["repository"], repo["url"]))
    _write_data_sheet(
        sheet, "GitHub Actions 工作流明细", report, headers, rows, "Workflows"
    )
    for row_number, url in enumerate(links, 4):
        _set_hyperlink(sheet.cell(row_number, 1), url)


def _write_metadata_sheet(sheet: Worksheet, report: Dict[str, Any]) -> None:
    period = report["period"]
    rows = [
        ["组织", report["organization"]],
        ["报告周期", period["label"]],
        ["本地开始", period["local_start"]],
        ["本地结束（不含）", period["local_end_exclusive"]],
        ["时区", period["timezone"]],
        ["生成时间", report["generated_at"]],
        ["数据源", report["coverage"]["source"]],
        ["认证方式", report["authentication"]],
        ["API Rate Limit", report["rate_limit"]],
    ]
    _write_title(sheet, "报告说明与指标口径", _period_description(period), 4)
    row = _write_summary_table(
        sheet, 4, "报告元数据", ["字段", "值"], rows, "ReportMetadata"
    )
    definitions = [
        ["默认分支 Commits", "统计期内各仓库默认分支提交数"],
        ["代码变更", "默认分支 Commit 的 additions + deletions"],
        ["首次响应", "统计期内创建的 Issue/PR 到首次非作者互动的时长"],
        ["处理时长", "统计期内创建且关闭的 Issue/PR 从创建到关闭的时长"],
        ["评审参与率", "收到至少一次 Review 的 PR 占比"],
        ["PR/Issue 关联率", "正文中关联 Issue 的 PR 占比"],
        ["Actions 总时长", "所有 Workflow Run 的 run_started_at 到 updated_at 墙钟时长之和"],
        ["首次通过率", "每个 Run 编号的首次 attempt 成功占比"],
        ["PR 覆盖率", "能关联到 Actions Run 的统计期 PR 占比"],
        ["Commit 覆盖率", "能关联到 Actions Run 的默认分支 Commit 占比"],
    ]
    row = _write_summary_table(
        sheet, row, "指标定义", ["指标", "定义"], definitions, "MetricDefinitions"
    )
    limitation_rows = [
        [index, limitation] for index, limitation in enumerate(report["limitations"], 1)
    ]
    _write_summary_table(
        sheet,
        row,
        "限制与注意事项",
        ["序号", "说明"],
        limitation_rows,
        "ReportLimitations",
    )
    sheet.freeze_panes = "A4"
    _finish_sheet(sheet, max_width=90)


def _write_data_sheet(
    sheet: Worksheet,
    title: str,
    report: Dict[str, Any],
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    table_name: str,
    percent_columns: Iterable[int] = (),
    decimal_columns: Iterable[int] = (),
) -> None:
    _write_title(sheet, title, _period_description(report["period"]), len(headers))
    for column, header in enumerate(headers, 1):
        sheet.cell(3, column, header)
    for row_number, values in enumerate(rows, 4):
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, _safe_value(value))
    if rows:
        _add_table(sheet, 3, 1, 3 + len(rows), len(headers), table_name)
    else:
        sheet.cell(4, 1, "本期无数据")
    for column in percent_columns:
        _format_column(sheet, column, 4, 3 + len(rows), PERCENT_FORMAT)
    for column in decimal_columns:
        _format_column(sheet, column, 4, 3 + len(rows), DECIMAL_FORMAT)
    sheet.freeze_panes = "A4"
    _finish_sheet(sheet)


def _write_summary_table(
    sheet: Worksheet,
    start_row: int,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    table_name: str,
    formats: Optional[Mapping[int, str]] = None,
) -> int:
    end_column = max(1, len(headers))
    sheet.merge_cells(
        start_row=start_row,
        start_column=1,
        end_row=start_row,
        end_column=end_column,
    )
    title_cell = sheet.cell(start_row, 1, title)
    title_cell.fill = SECTION_FILL
    title_cell.font = WHITE_FONT
    title_cell.alignment = Alignment(vertical="center")
    header_row = start_row + 1
    for column, header in enumerate(headers, 1):
        sheet.cell(header_row, column, header)
    first_data_row = header_row + 1
    for row_number, values in enumerate(rows, first_data_row):
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, _safe_value(value))
    if rows:
        _add_table(
            sheet,
            header_row,
            1,
            header_row + len(rows),
            len(headers),
            table_name,
        )
    if formats:
        for zero_based_column, number_format in formats.items():
            _format_column(
                sheet,
                zero_based_column + 1,
                first_data_row,
                header_row + len(rows),
                number_format,
            )
    return header_row + len(rows) + 2


def _write_title(
    sheet: Worksheet, title: str, subtitle: str, end_column: int
) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    title_cell = sheet.cell(1, 1, title)
    title_cell.fill = TITLE_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=16)
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    subtitle_cell = sheet.cell(2, 1, subtitle)
    subtitle_cell.fill = NOTE_FILL
    subtitle_cell.font = Font(color="4B5563", italic=True)
    subtitle_cell.alignment = Alignment(vertical="center")


def _add_table(
    sheet: Worksheet,
    start_row: int,
    start_column: int,
    end_row: int,
    end_column: int,
    display_name: str,
) -> None:
    reference = (
        f"{sheet.cell(start_row, start_column).coordinate}:"
        f"{sheet.cell(end_row, end_column).coordinate}"
    )
    table = Table(displayName=display_name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _add_repository_links(
    sheet: Worksheet, repositories: Sequence[Dict[str, Any]], column: int
) -> None:
    for row_number, repo in enumerate(repositories, 4):
        _set_hyperlink(sheet.cell(row_number, column), repo["url"])


def _set_hyperlink(cell: Cell, url: str) -> None:
    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")


def _format_column(
    sheet: Worksheet, column: int, first_row: int, last_row: int, number_format: str
) -> None:
    for row in range(first_row, last_row + 1):
        if isinstance(sheet.cell(row, column).value, (int, float)):
            sheet.cell(row, column).number_format = number_format


def _finish_sheet(sheet: Worksheet, max_width: int = 42) -> None:
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row >= 3 and not cell.fill.fill_type:
                cell.border = THIN_BORDER
    for column, column_cells in enumerate(sheet.columns, 1):
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        if not values:
            continue
        width = min(max(max(len(value), 8) for value in values) + 2, max_width)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.auto_filter.ref = None


def _period_description(period: Mapping[str, Any]) -> str:
    start = str(period["local_start"])
    end = str(period["local_end_exclusive"])
    try:
        inclusive_end = datetime.fromisoformat(end) - timedelta(microseconds=1)
        end_text = inclusive_end.date().isoformat()
    except ValueError:
        end_text = end
    return f"统计区间：{start[:10]} 至 {end_text}（{period['timezone']}）"


def _safe_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return str(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _hours_from_minutes(value: Optional[float]) -> Optional[float]:
    return None if value is None else value / 60


def _yes_no(value: bool) -> str:
    return "是" if value else "否"
