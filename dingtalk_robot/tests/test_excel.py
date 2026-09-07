import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from dingtalk_robot.oss_stats.excel import export_xlsx
from dingtalk_robot.oss_stats.metrics import collect_statistics
from dingtalk_robot.tests.test_metrics import FakeGitHubClient


class ExcelExportTest(unittest.TestCase):
    def test_exports_filterable_workbook(self):
        report = collect_statistics(
            FakeGitHubClient(),
            "OpenXiangShan",
            start=datetime(2026, 8, 1, tzinfo=timezone.utc),
            end=datetime(2026, 9, 1, tzinfo=timezone.utc),
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "2026-08.xlsx"
            export_xlsx(report, output)

            self.assertTrue(output.is_file())
            workbook = load_workbook(output, read_only=False, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "组织汇总",
                    "仓库活跃度",
                    "协作质量",
                    "Actions质量",
                    "CI覆盖",
                    "工作流明细",
                    "说明",
                ],
            )
            self.assertIn("RepositoryActivity", workbook["仓库活跃度"].tables)
            self.assertEqual(
                workbook["仓库活跃度"].cell(4, 2).hyperlink.target,
                "https://github.com/OpenXiangShan/demo",
            )
            self.assertEqual(workbook["组织汇总"].freeze_panes, "A4")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
